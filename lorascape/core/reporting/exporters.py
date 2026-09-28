# lorascape/core/reporting/exporters.py
"""
CoverageReport를 파일로 내보내는 함수들임. Qt에 의존하지 않음.
  - export_excel: 요약 / GW별 / GW별 단말 상세 / 단말별 / 미커버 단말 시트
  - export_nodes_csv: 단말별 표 하나 (CSV는 표 하나만 담을 수 있어서 가장 쓸모 있는 단말별을 씀)
  - build_report_html: PDF 인쇄용 HTML (실제 PDF 변환은 gui/report_pdf.py가 Qt로 함)
"""
from __future__ import annotations

import csv
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lorascape.core.reporting.coverage_report import (
    SF_LEVELS, RELATION_CONNECTED, REASON_LABELS,
)

XL_FONT = "Malgun Gothic"
GW_HEADERS = (["GW ID", "지역", "상세위치", "위도", "경도", "연결 단말", "수신 가능 단말", "다른 GW에도 수신"]
              + [f"SF{sf}" for sf in SF_LEVELS] + ["평균 Pr (dBm)"])
LINK_HEADERS = ["GW ID", "단말 ID", "유형", "지역", "상세위치", "구분", "거리 (km)", "Pr (dBm)", "SF", "SNR (dB)"]
NODE_HEADERS = ["단말 ID", "유형", "지역", "상세위치", "상태", "연결 GW", "Pr (dBm)", "SF", "수신 GW 수"]
UNCOVERED_HEADERS = ["단말 ID", "유형", "지역", "상세위치", "최대 수신 GW", "거리 (km)",
                     "최대 수신 (dBm)", "미커버 사유", "상세"]


def gw_table_rows(report):
    return [
        [g.gw_id, g.region, g.location_desc, g.lat, g.lon, g.connected, g.receivable, g.shared]
        + [g.sf_counts.get(sf, 0) for sf in SF_LEVELS] + [g.avg_rx_dbm]
        for g in report.gw_rows
    ]


def link_table_rows(report):
    return [
        [g.gw_id, l.node_id, l.device_type, l.region, l.location_desc, l.relation,
         l.distance_km, l.rx_power_dbm, l.sf, l.snr_db]
        for g in report.gw_rows for l in g.links
    ]


def node_table_rows(report):
    return [
        [n.node_id, n.device_type, n.region, n.location_desc, "연결" if n.covered else "미커버",
         n.gw_id or None, n.rx_power_dbm, n.sf, n.receiving_gws]
        for n in report.node_rows
    ]


def uncovered_table_rows(report):
    return [
        [u.node_id, u.device_type, u.region, u.location_desc, u.best_gw_id or None,
         u.distance_km, u.best_rx_dbm, u.reason_label, u.reason_detail]
        for u in report.uncovered_rows
    ]


def summary_pairs(report):
    """요약 시트/PDF 상단에 쓰는 (항목, 값) 목록임. 값은 화면 표시용 문자열임."""
    p = report.params
    pairs = [
        ("생성 시각", report.generated_at),
        ("전체 단말", str(report.total_nodes)),
        ("커버된 단말", str(report.covered_nodes)),
        ("미커버 단말", str(report.uncovered_nodes)),
        ("커버율", f"{report.coverage_ratio * 100:.1f}%"),
        ("목표 커버리지", f"{report.coverage_target * 100:.0f}% ({'달성' if report.target_met else '미달'})"),
        ("사용 GW", str(report.gw_count)),
        ("반송 주파수", f"{p.get('fc_mhz', 0):.1f} MHz"),
        ("환경 분류", str(p.get("environment", ""))),
        ("대역폭", f"{p.get('bandwidth_hz', 0):.0f} Hz"),
        ("경로손실 한계", f"{p.get('max_path_loss_db', 0):.0f} dB"),
    ]
    for code, cnt in sorted(report.reason_counts.items(), key=lambda kv: -kv[1]):
        pairs.append((f"미커버 사유 - {REASON_LABELS.get(code, code)}", str(cnt)))
    return pairs


# ── 엑셀 ────────────────────────────────────────────────────────────────

def _write_table(ws, headers, rows, widths, formats):
    head_fill = PatternFill("solid", start_color="DCE6F5")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=XL_FONT, bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=XL_FONT)
            if c in formats and value is not None:
                cell.number_format = formats[c]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def export_excel(report, path: str):
    """
    보고서를 엑셀로 저장함. 이 파일은 '검증 시점의 스냅샷'이라 수식 없이 값만 담음
    (원본 데이터가 바뀌어도 보고서가 자동으로 바뀌면 안 되기 때문).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.cell(row=1, column=1, value="GW 커버리지 보고서").font = Font(name=XL_FONT, bold=True, size=14)
    for i, (k, v) in enumerate(summary_pairs(report), start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=XL_FONT, bold=True)
        ws.cell(row=i, column=2, value=v).font = Font(name=XL_FONT)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34

    _write_table(wb.create_sheet("GW별"), GW_HEADERS, gw_table_rows(report),
                 [22, 14, 26, 11, 11, 11, 13, 16] + [7] * 6 + [14], {4: "0.000000", 5: "0.000000", 15: "0.0"})
    _write_table(wb.create_sheet("GW별 단말 상세"), LINK_HEADERS, link_table_rows(report),
                 [22, 24, 16, 14, 30, 9, 11, 11, 7, 10], {7: "0.00", 8: "0.0", 10: "0.0"})
    _write_table(wb.create_sheet("단말별"), NODE_HEADERS, node_table_rows(report),
                 [24, 16, 14, 30, 9, 22, 11, 7, 12], {7: "0.0"})
    _write_table(wb.create_sheet("미커버 단말"), UNCOVERED_HEADERS, uncovered_table_rows(report),
                 [24, 16, 14, 30, 22, 11, 15, 22, 44], {6: "0.00", 7: "0.0"})
    wb.save(path)


# ── CSV ─────────────────────────────────────────────────────────────────

def export_nodes_csv(report, path: str):
    """단말별 표를 CSV로 저장함 (utf-8-sig: 엑셀에서 바로 열어도 한글이 안 깨짐)."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(NODE_HEADERS)
        for row in node_table_rows(report):
            w.writerow(["" if v is None else v for v in row])


# ── PDF용 HTML ──────────────────────────────────────────────────────────

_TABLE_OPEN = ('<table width="100%" border="1" cellspacing="0" cellpadding="4" '
               'style="border-collapse:collapse; border-color:#b8c0d0">')


def _fmt(v, spec=None):
    if v is None:
        return "—"
    if spec:
        return format(v, spec)
    return escape(str(v))


def _html_table(headers, rows, aligns, specs=None):
    specs = specs or {}
    out = [_TABLE_OPEN, "<thead><tr>"]
    out += [f'<th bgcolor="#e6ecf7" align="center">{escape(h)}</th>' for h in headers]
    out.append("</tr></thead><tbody>")
    for row in rows:
        out.append("<tr>")
        for i, v in enumerate(row):
            out.append(f'<td align="{aligns[i]}">{_fmt(v, specs.get(i))}</td>')
        out.append("</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def build_report_html(report) -> str:
    """PDF 인쇄용 HTML을 만듦. 종이에 인쇄하기 좋게 흰 배경/어두운 글씨로 만듦."""
    parts = ['<html><body style="color:#1b2030">']
    parts.append("<h1>GW 커버리지 보고서</h1>")

    pairs = summary_pairs(report)
    parts.append(_TABLE_OPEN)
    for i in range(0, len(pairs), 2):
        parts.append("<tr>")
        for k, v in pairs[i:i + 2]:
            parts.append(f'<td bgcolor="#f1f4fa" width="20%"><b>{escape(k)}</b></td><td width="30%">{escape(v)}</td>')
        if len(pairs[i:i + 2]) == 1:
            parts.append('<td bgcolor="#f1f4fa"></td><td></td>')
        parts.append("</tr>")
    parts.append("</table>")

    parts.append("<h2>GW별 커버리지</h2>")
    parts.append("<p>연결 단말 = 그 GW에 실제로 붙는 단말, 수신 가능 = 그 GW의 신호가 닿는 단말(여러 GW에 겹쳐 셈)</p>")
    parts.append(_html_table(
        GW_HEADERS[:3] + GW_HEADERS[5:8] + GW_HEADERS[8:14] + [GW_HEADERS[14]],
        [[g.gw_id, g.region, g.location_desc, g.connected, g.receivable, g.shared]
         + [g.sf_counts.get(sf, 0) for sf in SF_LEVELS] + [g.avg_rx_dbm] for g in report.gw_rows],
        ["left", "left", "left", "right", "right", "right"] + ["right"] * 6 + ["right"],
        {12: ".1f"},
    ))

    parts.append('<h2 style="page-break-before:always">GW별 단말 상세</h2>')
    any_detail = False
    for g in report.gw_rows:
        if not g.links:
            continue
        any_detail = True
        types = ", ".join(f"{escape(t)} {c}" for t, c in g.type_counts.items()) or "—"
        parts.append(f"<h3>{escape(g.gw_id)} <small>({escape(g.region)} · {escape(g.location_desc)})</small></h3>")
        parts.append(f"<p>연결 {g.connected}개 · 수신 가능 {g.receivable}개 · 다른 GW에도 수신 {g.shared}개 · 연결 단말 유형: {types}</p>")
        parts.append(_html_table(
            LINK_HEADERS[1:],
            [[l.node_id, l.device_type, l.region, l.location_desc, l.relation, l.distance_km,
              l.rx_power_dbm, l.sf, l.snr_db] for l in g.links],
            ["left", "left", "left", "left", "center", "right", "right", "right", "right"],
            {5: ".2f", 6: ".1f", 8: ".1f"},
        ))
    if not any_detail:
        parts.append("<p>표시할 GW 연결 정보가 없습니다.</p>")

    parts.append('<h2 style="page-break-before:always">미커버 단말</h2>')
    if report.uncovered_rows:
        parts.append(_html_table(
            UNCOVERED_HEADERS,
            uncovered_table_rows(report),
            ["left", "left", "left", "left", "left", "right", "right", "left", "left"],
            {5: ".2f", 6: ".1f"},
        ))
    else:
        parts.append("<p>미커버 단말이 없습니다.</p>")

    parts.append("</body></html>")
    return "".join(parts)
