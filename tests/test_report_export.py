import csv
from openpyxl import load_workbook

from tests.test_coverage_report import gw, node, make_report
from lorascape.core.reporting.exporters import (
    export_excel, export_nodes_csv, build_report_html, NODE_HEADERS,
)


def sample_report():
    gws = [gw("A", 37.4000, 127.1200), gw("B<x>", 37.4002, 127.1202)]
    nodes = [node("N1", 37.4001, 127.1201, "맨홀수위센서"), node("N2", 38.50, 127.12, "CCTV")]
    rep, _ = make_report(gws, nodes)
    return rep


def test_excel_has_expected_sheets_and_rows(tmp_path):
    rep = sample_report()
    p = tmp_path / "r.xlsx"
    export_excel(rep, str(p))
    wb = load_workbook(p)
    assert wb.sheetnames == ["요약", "GW별", "GW별 단말 상세", "단말별", "미커버 단말"]
    assert wb["GW별"].max_row == 1 + len(rep.gw_rows)
    assert wb["단말별"].max_row == 1 + len(rep.node_rows)
    assert wb["미커버 단말"].max_row == 1 + len(rep.uncovered_rows)
    assert wb["GW별 단말 상세"].max_row == 1 + sum(len(g.links) for g in rep.gw_rows)
    assert wb["단말별"]["A1"].value == NODE_HEADERS[0]


def test_csv_roundtrip(tmp_path):
    rep = sample_report()
    p = tmp_path / "n.csv"
    export_nodes_csv(rep, str(p))
    with open(p, encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0] == NODE_HEADERS
    assert len(rows) == 1 + len(rep.node_rows)
    assert any(r[4] == "미커버" for r in rows[1:])


def test_html_escapes_and_contains_sections():
    html = build_report_html(sample_report())
    assert "B&lt;x&gt;" in html and "B<x>" not in html
    assert "GW별 커버리지" in html and "미커버 단말" in html and "경로손실 한계 초과" in html
